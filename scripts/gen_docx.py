#!/usr/bin/env python3
"""
生成 docx 文档：将多章正文合并为单个 Word 文件
用法：python3 gen_docx.py 章节文件1.md 章节文件2.md ... -o 输出路径.docx
"""

import sys
import re
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def parse_chapter(filepath):
    """读取 .md 文件，返回 (标题, 正文段落列表)"""
    with open(filepath, 'r', encoding='utf-8') as f:
        text = f.read()
    
    lines = text.strip().split('\n')
    
    # 第一行是标题
    title = lines[0].strip() if lines else "未命名章节"
    
    # 剩余是正文
    body_lines = [l.strip() for l in lines[1:] if l.strip()]
    
    return title, body_lines


def build_docx(chapters, output_path):
    """构建 docx，写入文件"""
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    # 设置段落间距
    paragraph_format = style.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(6)
    paragraph_format.line_spacing = 1.5
    
    for idx, (title, body_lines) in enumerate(chapters):
        if idx > 0:
            doc.add_page_break()
        
        # 章标题：居中、加粗、大字
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.add_run(title)
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.font.name = '黑体'
        title_run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        
        # 标题后空一行
        spacer = doc.add_paragraph()
        spacer_run = spacer.add_run('')
        spacer_run.font.size = Pt(6)
        
        # 正文段落
        for line in body_lines:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.first_line_indent = Cm(0.75)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(3)
            
            run = para.add_run(line)
            run.font.name = '宋体'
            run.font.size = Pt(11)
    
    doc.save(output_path)
    print(f"已生成: {output_path}")


def main():
    import argparse
    
    parser = argparse.ArgumentParser(description='将章节 md 文件合并为 docx')
    parser.add_argument('files', nargs='+', help='章节 .md 文件路径')
    parser.add_argument('-o', '--output', required=True, help='输出 .docx 路径')
    
    args = parser.parse_args()
    
    chapters = []
    for fp in args.files:
        if not os.path.exists(fp):
            print(f"文件不存在: {fp}")
            sys.exit(1)
        title, body = parse_chapter(fp)
        chapters.append((title, body))
        print(f"  {title} ({len(body)} 段)")
    
    build_docx(chapters, args.output)


if __name__ == '__main__':
    main()